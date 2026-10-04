"""results_table.py — Lưu kết quả và điền bảng thí nghiệm.

Nhiệm vụ: lưu kết quả từng lần chạy ra JSON, rồi điền vào experiments.xlsx từ mẫu
templates/experiment_table_template.xlsx (đừng gõ tay hàng chục dòng, rất dễ sai).

Tên cột của sheet "Experiments" (giữ nguyên, đúng thứ tự mẫu):
    exp_id, group, description, loss, optimizer, lr, weight_decay, batch, epochs, hidden, dropout,
    clip_norm, precision, init, seed, step0_loss, best_val_loss, best_epoch, final_train_loss,
    final_val_loss, val_acc, val_macro_f1, time_per_epoch_s, peak_mem_MB, diverged,
    eval_acc, eval_macro_f1, figure_file, notes
(các cột công thức ở cuối bảng mẫu tự tính, đừng ghi đè)
"""
from __future__ import annotations

import json
from pathlib import Path

GROUP_ALIASES = {"learning_rate": "hparam", "initialization": "init"}


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi result["cfg"], result["history"], result["summary"] (KHÔNG ghi best_state) ra
    <results_dir>/<exp_id>.json. Trả về đường dẫn file. Tạo thư mục nếu chưa có."""
    if "cfg" not in result or "history" not in result or "summary" not in result:
        raise ValueError("result phải chứa cfg, history và summary")
    exp_id = result["cfg"].get("exp_id")
    if not exp_id:
        raise ValueError("result['cfg']['exp_id'] không được rỗng")
    directory = Path(results_dir)
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / f"{exp_id}.json"
    serializable = {key: result[key] for key in ("cfg", "history", "summary")}
    with path.open("w", encoding="utf-8") as handle:
        json.dump(serializable, handle, ensure_ascii=False, indent=2, allow_nan=False)
    return str(path)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, trả về danh sách dict (sắp theo exp_id)."""
    directory = Path(results_dir)
    if not directory.exists():
        return []
    results = []
    for path in sorted(directory.glob("*.json")):
        with path.open(encoding="utf-8") as handle:
            result = json.load(handle)
        if not {"cfg", "history", "summary"}.issubset(result):
            raise ValueError(f"{path} không phải file kết quả hợp lệ")
        results.append(result)
    return sorted(results, key=lambda result: result["cfg"].get("exp_id", ""))


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Biến một kết quả thành một dòng của bảng: gộp cfg + summary (+ eval_acc, eval_macro_f1 nếu có)
    + figure_file = f"figures/{exp_id}.png". Khoá phải trùng tên cột ở đầu file.
    Chỉ truyền eval_scores cho baseline và cấu hình cuối cùng."""
    cfg, summary = result["cfg"], result["summary"]
    exp_id = cfg["exp_id"]
    row = {
        "exp_id": exp_id,
        "group": GROUP_ALIASES.get(cfg.get("group", ""), cfg.get("group", "")),
        "description": cfg.get("description", ""),
        "loss": cfg.get("loss", ""),
        "optimizer": cfg.get("optimizer", ""),
        "lr": cfg.get("lr"),
        "weight_decay": cfg.get("weight_decay"),
        "batch": cfg.get("batch"),
        "epochs": cfg.get("epochs"),
        "hidden": str(tuple(cfg.get("hidden", ()))),
        "dropout": cfg.get("dropout"),
        "clip_norm": cfg.get("clip_norm"),
        "precision": cfg.get("precision", ""),
        "init": cfg.get("init", ""),
        "seed": cfg.get("seed"),
        **{key: summary.get(key) for key in (
            "step0_loss", "best_val_loss", "best_epoch", "final_train_loss", "final_val_loss",
            "val_acc", "val_macro_f1", "time_per_epoch_s", "peak_mem_MB", "diverged",
        )},
        "eval_acc": None if eval_scores is None else eval_scores.get("accuracy", eval_scores.get("eval_acc")),
        "eval_macro_f1": None if eval_scores is None else eval_scores.get("macro_f1", eval_scores.get("eval_macro_f1")),
        "figure_file": f"figures/{exp_id}.png",
        "notes": notes,
    }
    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str) -> None:
    """Điền các dòng vào sheet "Experiments" của mẫu, từ dòng 2 trở xuống, rồi lưu thành out_path.

    Các bước (openpyxl):
      1. wb = openpyxl.load_workbook(template_path)   # KHÔNG dùng data_only=True (sẽ mất công thức)
      2. ws = wb["Experiments"]; đọc tiêu đề dòng 1 để biết cột nào ứng với khoá nào
      3. với mỗi row: ghi giá trị vào đúng cột; BỎ QUA các cột công thức (step0_gap_vs_lnC, gap_val_minus_train,
         delta_val_f1_vs_base, beyond_noise)
      4. wb.save(out_path)
    Sau khi lưu, mở file bằng Excel/LibreOffice để các công thức tính lại.
    """
    from openpyxl import load_workbook

    workbook = load_workbook(template_path)
    worksheet = workbook["Experiments"]
    headers = [cell.value for cell in worksheet[1]]
    formula_columns = {"step0_gap_vs_lnC", "gap_val_minus_train", "delta_val_f1_vs_base", "beyond_noise"}
    for row_number, row in enumerate(rows, start=2):
        for column_number, header in enumerate(headers, start=1):
            if header in formula_columns or header not in row:
                continue
            value = row[header]
            if header == "group":
                value = GROUP_ALIASES.get(value, value)
            worksheet.cell(row=row_number, column=column_number).value = value

    # Only baseline IDs are inputs. Preserve lookup and mean/std/2-sigma formulas.
    baseline_ids = [row["exp_id"] for row in rows if row.get("group") == "baseline"]
    if len(baseline_ids) > 5:
        raise ValueError("Sheet Seeds của template hỗ trợ tối đa 5 baseline seed")
    seed_sheet = workbook["Seeds"]
    for index in range(2, 7):
        seed_sheet.cell(index, 1).value = baseline_ids[index - 2] if index - 2 < len(baseline_ids) else None

    # Keep Summary group labels/formulas and recalculate when opened in Excel.
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    destination = Path(out_path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
